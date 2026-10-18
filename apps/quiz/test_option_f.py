from io import BytesIO, StringIO
from tempfile import NamedTemporaryFile

from django.contrib.auth import get_user_model
from django.core.management import call_command
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import Student
from apps.api.views import _generate_jwt
from apps.bot.extractor import validate_questions
from apps.content.models import Course, Department
from .importers import import_exit_questions_from_excel, import_questions_from_excel
from .models import ExamPaper, Question, QuizSimulation


class SixthChoiceTests(APITestCase):
    def setUp(self):
        self.course = Course.objects.create(name='Six choices', code='SIX101')
        self.paper = ExamPaper.objects.create(title='Six-choice quiz', course=self.course,
            exam_type='quiz', year=2026, duration_minutes=30, access_level='free')
        department = Department.objects.create(name='Management')
        self.exit_paper = ExamPaper.objects.create(title='Exit model', department=department,
            exam_type=ExamPaper.TYPE_EXIT_MODEL, year=2026, duration_minutes=30, access_level='free')
        self.student = Student.objects.create(telegram_id=601234, first_name='Six-choice learner')
        self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + _generate_jwt(self.student))

    def question(self, **changes):
        return Question.objects.create(exam_paper=self.paper, text='Choose the sixth answer.',
            **{**{f'option_{letter}': f'Choice {letter.upper()}' for letter in 'abcdef'},
               'correct_option': 'f', 'topic_tags': ['Chapter 1'], **changes})

    def workbook(self, *, include_f=True, **changes):
        headers = ['question_type', 'question', 'chapter', 'option_a', 'option_b',
                   'option_c', 'option_d', 'option_e', 'correct_option', 'explanation']
        if include_f:
            headers.insert(8, 'option_f')
        row = {'question_type': 'mcq', 'question': 'Choose the sixth answer.', 'chapter': 'Chapter 1',
               **{f'option_{letter}': f'Choice {letter.upper()}' for letter in 'abcdef'},
               'correct_option': 'F', 'explanation': 'The sixth choice is correct.', **changes}
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Questions'
        sheet.append(headers)
        sheet.append(['Full question text' if header == 'question' else '' for header in headers])
        sheet.append([row.get(header, '') for header in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return output

    def submit(self, question, answer='f', **changes):
        return self.client.post('/api/quiz/attempts/', {
            'exam_paper_id': question.exam_paper_id,
            'answers': [{'question_id': question.pk, 'selected_option': answer}], **changes,
        }, format='json')

    def test_practice_exposes_six_choices_and_grades_f(self):
        question = self.question()
        response = self.client.get(f'/api/exams/{self.paper.pk}/questions/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]['option_f'], 'Choice F')
        self.assertEqual(list(response.data[0]['available_options']), list('abcdef'))
        response = self.submit(question, answer='F')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['score'], 1)
        detail = response.data['detailed_answers'][str(question.pk)]
        self.assertTrue(detail['is_correct'])
        self.assertEqual(detail['options']['f'], 'Choice F')

    def test_simulation_preserves_issued_sixth_choice_during_later_edits(self):
        question = self.question()
        response = self.client.get(f'/api/exams/{self.paper.pk}/questions/', {'mode': 'simulation'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]['available_options']['f'], 'Choice F')
        self.assertNotIn('correct_option', response.data[0])
        simulation_id = response.headers['X-Quiz-Simulation-ID']
        Question.objects.filter(pk=question.pk).update(option_f='Changed after issue', correct_option='a')
        response = self.submit(question, mode='simulation', simulation_id=simulation_id)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['score'], 1)
        self.assertEqual(response.data['detailed_answers'][str(question.pk)]['options']['f'], 'Choice F')

    def test_simulations_issued_before_option_f_still_resume_and_submit(self):
        question = self.question(option_e='', option_f='', correct_option='a')
        response = self.client.get(f'/api/exams/{self.paper.pk}/questions/', {'mode': 'simulation'})
        simulation = QuizSimulation.objects.get(pk=response.headers['X-Quiz-Simulation-ID'])
        for row in simulation.question_snapshot:
            del row['option_f']
        simulation.save(update_fields=['question_snapshot'])
        response = self.client.get(f'/api/exams/{self.paper.pk}/questions/', {'mode': 'simulation'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['X-Quiz-Simulation-ID'], str(simulation.pk))
        response = self.submit(question, answer='a', mode='simulation', simulation_id=str(simulation.pk))
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['score'], 1)

    def test_f_cannot_be_submitted_for_questions_without_that_choice(self):
        for option_e, answer in (('', 'a'), ('Choice E', 'e')):
            with self.subTest(option_e=option_e):
                question = self.question(option_e=option_e, option_f='', correct_option=answer)
                self.assertNotIn('f', question.available_options)
                self.assertEqual(self.submit(question).status_code, 400)
                response = self.submit(question, answer=answer)
                self.assertEqual(response.status_code, 201)
                self.assertEqual(response.data['score'], 1)

    def test_selective_practice_can_display_and_score_f(self):
        question = self.question()
        response = self.client.post('/api/quiz/selective-practice/', {
            'course_id': self.course.pk, 'selected_topics': ['Chapter 1'], 'limit': 10,
        }, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['questions'][0]['available_options']['f'], 'Choice F')
        response = self.submit(question, mode='selective', course_id=self.course.pk, selected_topics=['Chapter 1'])
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['score'], 1)

    def test_both_excel_importers_accept_f_and_reject_missing_sixth_choice_text(self):
        for importer, paper in ((import_questions_from_excel, self.paper),
                                (import_exit_questions_from_excel, self.exit_paper)):
            with self.subTest(importer=importer.__name__):
                result = importer(self.workbook(), paper)
                self.assertEqual((result.created, result.skipped), (1, 0))
                question = paper.questions.get()
                self.assertEqual((question.option_f, question.correct_option, question.is_active), ('Choice F', 'f', False))
                for workbook in (self.workbook(option_f=''), self.workbook(include_f=False)):
                    result = importer(workbook, paper)
                    self.assertEqual((result.created, result.skipped), (0, 1))
                    self.assertIn('option_f is required', result.errors[0])
                result = importer(self.workbook(include_f=False, correct_option='e'), paper)
                self.assertEqual((result.created, result.skipped), (1, 0))

    def test_true_false_f_is_rejected_and_extra_choices_are_not_exposed(self):
        result = import_exit_questions_from_excel(self.workbook(question_type='true_false'), self.exit_paper)
        self.assertEqual((result.created, result.skipped), (0, 1))
        question = self.question(question_type='true_false', correct_option='b', option_a='True', option_b='False')
        self.assertEqual(question.available_options, {'a': 'True', 'b': 'False'})

    def test_course_chapter_command_preserves_sixth_choice(self):
        with NamedTemporaryFile(suffix='.xlsx') as file:
            file.write(self.workbook().read())
            file.flush()
            call_command('import_course_with_chapters', file=file.name, course_id=self.course.pk,
                         exam_paper_id=self.paper.pk, stdout=StringIO())
        question = self.paper.questions.get()
        self.assertEqual((question.option_f, question.correct_option), ('Choice F', 'f'))

    def test_admin_form_and_both_downloaded_templates_support_option_f(self):
        self.client.force_login(get_user_model().objects.create_superuser('six-choice-admin'))
        response = self.client.get('/admin/quiz/question/add/')
        self.assertContains(response, 'name="option_f"')
        response = self.client.post('/admin/quiz/question/add/', {
            'text': 'Created in admin with F', 'question_type': 'mcq', 'exam_paper': self.paper.pk,
            **{f'option_{letter}': f'Choice {letter.upper()}' for letter in 'abcdef'},
            'correct_option': 'f', 'difficulty': 'medium', 'topic_tags': '[]', 'is_active': 'on',
        })
        self.assertEqual(response.status_code, 302)
        question = self.paper.questions.get()
        self.assertEqual((question.option_f, question.correct_option), ('Choice F', 'f'))
        for paper in (self.paper, self.exit_paper):
            with self.subTest(paper=paper):
                response = self.client.get(f'/admin/quiz/exampaper/{paper.pk}/download-template/')
                self.assertEqual(response.status_code, 200)
                workbook = load_workbook(BytesIO(response.content))
                headers = [cell.value for cell in workbook['Questions'][1]]
                self.assertIn('option_f', headers)
                workbook.close()

    def test_ai_validation_preserves_f_and_rejects_answer_without_option_text(self):
        row = {'question_type': 'mcq', 'question': 'Choose F', 'correct_option': 'F',
               **{f'option_{letter}': f'Choice {letter.upper()}' for letter in 'abcdef'}}
        result = validate_questions([row])[0]
        self.assertEqual((result['option_f'], result['correct_option']), ('Choice F', 'f'))
        for changes in ({'option_f': ''}, {'question_type': 'true_false'}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_questions([{**row, **changes}])
