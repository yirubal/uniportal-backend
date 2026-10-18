from io import BytesIO, StringIO
from tempfile import NamedTemporaryFile

from django.contrib.auth import get_user_model
from django.core.management import call_command, CommandError
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.content.models import Course, Department
from .importers import import_exit_questions_from_excel
from .models import Chapter, ExamPaper, Question


class ExitQuestionImportTests(TestCase):
    def setUp(self):
        self.department = Department.objects.create(name='Management')
        self.paper = ExamPaper.objects.create(title='Management Model', department=self.department,
            exam_type=ExamPaper.TYPE_EXIT_MODEL, year=2026)

    def workbook(self, rows, headers=None, notes=True):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Questions'
        headers = headers or ['course', 'question', 'question_type', 'option_a',
                              'option_b', 'correct_option', 'topic_tags']
        sheet.append(headers)
        if notes:
            sheet.append(['Full question text' if name == 'question' else None for name in headers])
        for row in rows:
            sheet.append([row.get(name) for name in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return output

    def row(self, **overrides):
        return {'question': 'Which leadership style involves the team?', 'question_type': 'mcq',
                'option_a': 'Participative', 'option_b': 'Autocratic', 'correct_option': 'a', **overrides}

    def test_official_and_model_questions_need_neither_course_nor_chapter(self):
        for exam_type in (ExamPaper.TYPE_EXIT_REAL, ExamPaper.TYPE_EXIT_MODEL):
            with self.subTest(exam_type=exam_type):
                self.paper.exam_type = exam_type
                self.paper.save(update_fields=['exam_type'])
                result = import_exit_questions_from_excel(self.workbook([self.row(course='  ')]), self.paper)
                self.assertEqual((result.created, result.skipped, result.errors), (1, 0, []))
        self.paper.refresh_from_db()
        self.assertIsNone(self.paper.course)
        self.assertFalse(Course.objects.exists())
        self.assertFalse(Chapter.objects.exists())
        self.assertEqual(self.paper.questions.filter(chapter__isnull=True, is_active=False).count(), 2)

    def test_course_column_and_notes_row_can_be_omitted(self):
        headers = ['question', 'question_type', 'option_a', 'option_b', 'correct_option', 'topic_tags']
        result = import_exit_questions_from_excel(self.workbook([
            self.row(topic_tags='leadership,ethics,leadership'),
            self.row(question='Leaders can listen.', question_type='true_false',
                     option_a='', option_b='', correct_option='true'),
            self.row(question='Discuss ethical leadership.', question_type='essay', correct_option=''),
        ], headers=headers, notes=False), self.paper)
        self.assertEqual((result.created, result.skipped, result.errors), (3, 0, []))
        self.assertEqual(self.paper.questions.get(question_type='mcq').topic_tags, ['leadership', 'ethics'])
        true_false = self.paper.questions.get(question_type='true_false')
        self.assertEqual((true_false.option_a, true_false.option_b, true_false.correct_option), ('True', 'False', 'a'))

    def test_blank_course_at_row_167_is_imported(self):
        output = self.workbook([])
        workbook = load_workbook(output)
        for column, value in enumerate(self.row().values(), start=2):
            workbook['Questions'].cell(row=167, column=column, value=value)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        result = import_exit_questions_from_excel(output, self.paper)
        self.assertEqual((result.created, result.skipped, result.errors), (1, 0, []))

    def test_optional_course_keeps_existing_paper_assignment_and_validates_supplied_values(self):
        course = Course.objects.create(name='Leadership', code='LEAD101')
        other = Course.objects.create(name='Accounting', code='ACC101')
        self.paper.course = course
        self.paper.save(update_fields=['course'])
        result = import_exit_questions_from_excel(self.workbook([
            self.row(), self.row(course=course.code),
            self.row(course=other.code), self.row(course='Unknown course'),
        ]), self.paper)
        self.assertEqual((result.created, result.skipped), (2, 2))
        self.assertIn('does not match', result.errors[0])
        self.assertIn('was not found', result.errors[1])
        self.paper.refresh_from_db()
        self.assertEqual(self.paper.course_id, course.pk)

    def test_missing_text_invalid_type_and_invalid_answer_still_report_row_errors(self):
        result = import_exit_questions_from_excel(self.workbook([
            self.row(question=''), self.row(question_type='invalid'), self.row(correct_option='invalid'),
            self.row(question='A valid question'),
        ]), self.paper)
        self.assertEqual((result.created, result.skipped), (1, 3))
        self.assertTrue(result.errors[0].startswith('Row 3:'))
        self.assertIn('question text is empty', result.errors[0])
        self.assertIn('unknown question_type', result.errors[1])
        self.assertIn('invalid correct_option', result.errors[2])

    def test_failed_command_replacement_preserves_existing_questions(self):
        original = Question.objects.create(exam_paper=self.paper, text='Keep this question')
        with NamedTemporaryFile(suffix='.xlsx') as file:
            file.write(self.workbook([self.row(correct_option='invalid')]).read())
            file.flush()
            with self.assertRaises(CommandError):
                call_command('import_exit_exam_questions', file=file.name,
                             exam_paper_id=self.paper.pk, delete_existing=True, stdout=StringIO())
        self.assertTrue(Question.objects.filter(pk=original.pk).exists())

    def test_command_supports_named_questions_sheet_without_course_or_chapter(self):
        workbook = load_workbook(self.workbook([self.row()], notes=False))
        sheet = workbook['Questions']
        sheet.title = 'Import data'
        sheet.title = 'QUESTIONS'
        workbook.create_sheet('Instructions', 0)
        workbook.active = 0
        with NamedTemporaryFile(suffix='.xlsx') as file:
            workbook.save(file.name)
            workbook.close()
            call_command('import_exit_exam_questions', file=file.name,
                         exam_paper_id=self.paper.pk, stdout=StringIO())
        self.assertTrue(self.paper.questions.get().is_active)
        self.paper.refresh_from_db()
        self.assertIsNone(self.paper.course)

    def test_admin_download_and_upload_of_official_template_with_blank_course(self):
        user = get_user_model().objects.create_superuser('exit-import-admin')
        self.client.force_login(user)
        base = f'/admin/quiz/exampaper/{self.paper.pk}/'
        response = self.client.get(base + 'import-questions/')
        self.assertContains(response, 'Course and chapter are not required.')
        response = self.client.get(base + 'download-template/')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook['Questions']
        self.assertTrue(sheet['A2'].value.startswith('Optional.'))
        row = {'question': 'What is ethical leadership?', 'option_a': 'Accountability',
               'option_b': 'Deception', 'option_c': 'Favoritism', 'option_d': 'Silence',
               'correct_option': 'a', 'difficulty': 'medium', 'topic_tags': 'ethics'}
        sheet.append([row.get(cell.value, '') for cell in sheet[1]])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        output.name = 'exit-questions.xlsx'
        response = self.client.post(base + 'import-questions/', {'excel_file': output})
        self.assertEqual(response.status_code, 302)
        question = self.paper.questions.get()
        self.assertEqual(question.topic_tags, ['ethics'])
        self.assertFalse(question.is_active)
        self.assertIsNone(question.chapter)
        self.paper.refresh_from_db()
        self.assertIsNone(self.paper.course)
